import pdfplumber
import re
from openpyxl import Workbook
import fitz  # PyMuPDF

# Chargement du PDF

pdf_path = "PV CREDITFLOW (3).pdf"
def add_cropbox_if_missing(pdf_path, output_path):
	doc = fitz.open(pdf_path)

	for page in doc:
		mediabox = page.rect
		cropbox = page.cropbox
		
		if cropbox == mediabox:
			# Réduire manuellement les bords (par exemple 10 pts de chaque côté)
			new_cropbox = fitz.Rect(
				mediabox.x0 + 10,
				mediabox.y0 + 10,
				mediabox.x1 - 10,
				mediabox.y1 - 10
			)
			page.set_cropbox(new_cropbox)

	doc.save(output_path)
	print(f"PDF corrigé : {output_path}")

# Utilisation
final_path = "final.pdf"
add_cropbox_if_missing(pdf_path, final_path)

pdf_path = final_path

# Dictionnaire pour stocker les données
data = {
	"Numéro de PV": None,
	"Emprunteur": None,
	"CAF": None,
	"Nature juridique": None,
	"N° de Compte / Matricule": None,
	"Activité principale": None,
	"Objet du financement": None,
	"Type de concours sollicité": None,
	"Montant sollicité": None,
	"Montant proposé par le CAF": None,
	"Durée proposée (ARC)": None,
	"Garanties": None,
	"Chiffre d'affaires mensuel": None,
	"Marge totale": None,
	"Capacité dégagée": None,
	"TMC": None
}

def normalize_keywords(text):
	# Liste des mots-clés susceptibles d'être coupés
	fixes = [
		("Nature\njuridique", "Nature juridique"),
		("Type de\nconcours", "Type de concours"),
		("Montant\nsollicité", "Montant sollicité"),
		("Montant\nproposé par le\nCAF", "Montant proposé par le CAF"),
		("Activité\nprincipale", "Activité principale"),
		("Objet du\nfinancement", "Objet du financement"),
		("N° de Compte\n/ Matricule", "N° de Compte / Matricule"),
	]
	for broken, fixed in fixes:
		text = text.replace(broken, fixed)
	return text


# Fonctions d’extraction avec regex
def extract_value(pattern, text, group=1):
	match = re.search(pattern, text, re.IGNORECASE)
	return match.group(group).strip() if match else None

# Extraction de texte
with pdfplumber.open(pdf_path) as pdf:
	full_text = "\n".join(page.extract_text() for page in pdf.pages if page.extract_text())
	full_text = normalize_keywords(full_text)
	print(full_text)
	data["Numéro de PV"] = extract_value(r'PROCÈS VERBAL\s*-\s*(CFNTG-\S+)', full_text)
	data["Emprunteur"] = extract_value(r'Emprunteur\s+(.*?)\n', full_text)
	data["CAF"] = extract_value(r'CAF\s+(.*?)\n', full_text)
	data["Nature juridique"] = extract_value(r'Nature\s+juridique\s+(.*?)\n', full_text)
	data["N° de Compte / Matricule"] = extract_value(r'N° de Compte\s*/\s*Matricule\s+(.*?)\n', full_text)
	data["Activité principale"] = extract_value(r'Activité\s+principale\s+(.*?)\n', full_text)
	data["Objet du financement"] = extract_value(r'Objet du\s+financement\s+(.*?)\n', full_text)
	data["Type de concours sollicité"] = extract_value(r'Type de\s+concours\s+sollicité\s+(.*?)\n', full_text)
	data["Montant sollicité"] = extract_value(r'Montant\s+sollicité\s+([0-9\s]+)', full_text)
	data["Montant proposé par le CAF"] = extract_value(r'Montant\s+proposé par le\s+CAF\s+([0-9\s]+)', full_text)
	data["Durée proposée (ARC)"] = extract_value(r'Durée\s+\d+\s+\d+\s+(\d+)', full_text)
	data["Chiffre d'affaires mensuel"] = extract_value(r'un chiffre d\'affaires mensuel.*?([0-9\s]+M)', full_text)
	data["Marge totale"] = extract_value(r'marge totale.*?([0-9\s]+)', full_text)
	data["Capacité dégagée"] = extract_value(r'capacité dégagée.*?XOF\s+([0-9\s]+)', full_text)
	data["TMC"] = extract_value(r'TMC confié.*?FCFA\s+([0-9.,\s]+)', full_text)

	# Garanties simplifiées
	if "Déposit de 20%" in full_text and "Cautions Personnelle et Solidaire" in full_text:
		data["Garanties"] = "Déposit 20% + 2 cautions solidaires"

# Création du fichier Excel
wb = Workbook()
ws = wb.active
ws.title = "Données du PV"

# Remplir les données
for idx, (key, value) in enumerate(data.items(), start=1):
	ws.cell(row=idx, column=1, value=key)
	ws.cell(row=idx, column=2, value=value)

# Enregistrement
wb.save("donnees_pv_creditflow.xlsx")
print("✅ Données extraites et sauvegardées dans 'donnees_pv_creditflow.xlsx'")
